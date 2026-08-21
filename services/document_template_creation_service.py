from __future__ import annotations

import hashlib
import os
import re
import shutil
import unicodedata
import uuid
from pathlib import Path, PurePosixPath
from typing import Any, BinaryIO, Iterable

from docx import Document

from services.ai_planner_document_service import validate_ai_planner_template_structure
from services.cross_process_file_lock import CrossProcessFileLock
from services.document_template_generation_service import generate_synthetic_document
from services.document_template_read_service import (
    ResolvedDocument,
    build_document_whitelist,
    clear_document_template_read_cache,
)
from services.document_template_runtime_service import atomic_write_bytes, atomic_write_json, read_json
from services.document_template_storage_service import (
    MAX_UPLOAD_BYTES,
    CandidateUploadTooLarge,
    cache_root,
    candidate_directory,
    data_root,
    list_candidates,
    list_history,
    utc_now,
    validate_uuid,
)
from services.document_template_validation_service import (
    LEGACY_PLACEHOLDERS,
    ValidationFailure,
    build_contract,
    inspect_docx,
)
from services.release_template_catalog_service import (
    build_runtime_template_catalog,
    clear_template_catalog_cache,
)


KIT_ID_RE = re.compile(r"^dtk1_[0-9a-f]{64}$")
KE_RE = re.compile(r"^[0-9]{5,20}$")
WINDOWS_RESERVED = {
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{number}" for number in range(1, 10)),
    *(f"LPT{number}" for number in range(1, 10)),
}


def _placeholder_name(value: str) -> str:
    text = str(value or "").strip()
    wrappers = (("{{", "}}"), ("${", "}"), ("<<", ">>"), ("[[", "]]"))
    for prefix, suffix in wrappers:
        if text.startswith(prefix) and text.endswith(suffix):
            return text[len(prefix):-len(suffix)].strip()
    return text


class CreationDraftNotFound(LookupError):
    pass


class CreationDraftConflict(RuntimeError):
    pass


def _normalized_component(value: str) -> str:
    value = unicodedata.normalize("NFC", str(value or "").strip())
    if (
        not value
        or value in {".", ".."}
        or value[-1:] in {" ", "."}
        or any(ord(char) < 32 for char in value)
        or any(char in '<>:"/\\|?*' for char in value)
        or value.split(".", 1)[0].upper() in WINDOWS_RESERVED
    ):
        raise ValueError("Имя содержит недопустимые символы.")
    return value


def safe_docx_filename(value: str) -> str:
    value = _normalized_component(Path(str(value or "")).name)
    if len(value) > 180 or not value.casefold().endswith(".docx"):
        raise ValueError("Укажите безопасное имя DOCX длиной до 180 символов.")
    return value


def safe_kit_name(value: str) -> str:
    value = _normalized_component(value)
    if len(value) > 120 or re.search(r"\([0-9]{5,20}\)\s*$", value):
        raise ValueError("Название комплекта должно быть без КЭ в скобках и не длиннее 120 символов.")
    return value


def kit_id_for_relative_dir(relative_dir: str) -> str:
    normalized = unicodedata.normalize("NFC", str(PurePosixPath(str(relative_dir or "")))).replace("\\", "/")
    parsed = PurePosixPath(normalized)
    if parsed.is_absolute() or len(parsed.parts) != 2 or any(part in {"", ".", ".."} for part in parsed.parts):
        raise ValueError("Invalid kit path")
    digest = hashlib.sha256(parsed.as_posix().casefold().encode("utf-8")).hexdigest()
    return f"dtk1_{digest}"


def list_kit_models(root: Path) -> list[dict[str, Any]]:
    whitelist = build_document_whitelist(Path(root), fresh=True)
    models = []
    for entry in build_runtime_template_catalog(root=Path(root)):
        relative_dir = str(entry.get("relative_dir") or "")
        try:
            kit_id = kit_id_for_relative_dir(relative_dir)
        except ValueError:
            continue
        filenames = sorted(
            item.filename
            for item in whitelist.values()
            if item.relative_directory.casefold() == relative_dir.casefold()
        )
        models.append({**entry, "kit_id": kit_id, "filenames": filenames})
    return models


def resolve_kit(kit_id: str, root: Path) -> dict[str, Any] | None:
    if not KIT_ID_RE.fullmatch(str(kit_id or "")):
        return None
    for model in list_kit_models(root):
        if model["kit_id"] == kit_id:
            directory = Path(root) / PurePosixPath(model["relative_dir"])
            try:
                if directory.is_symlink() or not directory.resolve(strict=True).is_relative_to(Path(root).resolve(strict=True)):
                    return None
            except (OSError, RuntimeError, ValueError):
                return None
            return {**model, "path": directory}
    return None


def _draft_directory(draft_uuid: str) -> Path:
    return cache_root() / "creation-drafts" / validate_uuid(draft_uuid)


def _draft_lock(draft_uuid: str) -> CrossProcessFileLock:
    value = validate_uuid(draft_uuid)
    return CrossProcessFileLock(cache_root() / "locks" / f"creation-{value}.lock", timeout=5)


def _kit_lock(kit_id: str) -> CrossProcessFileLock:
    if not KIT_ID_RE.fullmatch(str(kit_id or "")):
        raise ValueError("Invalid kit identifier")
    return CrossProcessFileLock(cache_root() / "locks" / f"kit-{kit_id}.lock", timeout=5)


def _write_stream(stream: BinaryIO, target: Path) -> tuple[int, str]:
    size = 0
    digest = hashlib.sha256()
    with target.open("xb") as output:
        while True:
            chunk = stream.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > MAX_UPLOAD_BYTES:
                raise CandidateUploadTooLarge("File exceeds the 10 MiB limit")
            output.write(chunk)
            digest.update(chunk)
        output.flush()
        os.fsync(output.fileno())
    if not size:
        raise ValueError("Выберите непустой DOCX-файл.")
    return size, digest.hexdigest()


def _ensure_unique_names(names: Iterable[str]) -> list[str]:
    safe_names = [safe_docx_filename(name) for name in names]
    normalized = [unicodedata.normalize("NFC", name).casefold() for name in safe_names]
    if len(normalized) != len(set(normalized)):
        raise ValueError("Имена документов внутри комплекта должны быть уникальными.")
    return safe_names


def _new_draft(metadata: dict[str, Any]) -> tuple[str, Path]:
    draft_uuid = str(uuid.uuid4())
    directory = _draft_directory(draft_uuid)
    directory.mkdir(parents=True, exist_ok=False)
    metadata.update({
        "version": 1,
        "draft_uuid": draft_uuid,
        "state": "uploaded",
        "created_at": utc_now(),
        "updated_at": utc_now(),
        "validation": None,
    })
    atomic_write_json(directory / "metadata.json", metadata)
    return draft_uuid, directory


def create_kit_draft(
    *, root: Path, category: str, name: str, ke: str, variant: str = "", aliases: list[str] | None = None,
    source_mode: str, uploads: list[tuple[str, BinaryIO]] | None = None, source_kit_id: str = "",
    target_names: list[str] | None = None,
) -> dict[str, Any]:
    root = Path(root)
    categories = {str(entry.get("category") or "") for entry in build_runtime_template_catalog(root=root)}
    if category not in categories:
        raise ValueError("Выберите существующую категорию.")
    ke = str(ke or "").strip()
    if not KE_RE.fullmatch(ke):
        raise ValueError("КЭ должен содержать от 5 до 20 цифр.")
    variant = unicodedata.normalize("NFC", str(variant or "").strip())[:80]
    aliases = [unicodedata.normalize("NFC", str(value or "").strip())[:120] for value in aliases or []]
    aliases = [value for value in aliases if value]
    duplicates = [entry for entry in build_runtime_template_catalog(root=root) if entry.get("category") == category and str(entry.get("ke") or "") == ke]
    if duplicates and not (variant or aliases):
        raise ValueError("Этот КЭ уже используется. Укажите отличительный вариант или алиас.")
    upload_items = list(uploads or [])
    source_kit = None
    source_docs: list[ResolvedDocument] = []
    if source_mode == "copy":
        source_kit = resolve_kit(source_kit_id, root)
        if source_kit is None:
            raise ValueError("Комплект-основа не найден.")
        whitelist = build_document_whitelist(root, fresh=True)
        source_docs = sorted(
            [item for item in whitelist.values() if item.relative_directory.casefold() == str(source_kit["relative_dir"]).casefold()],
            key=lambda item: item.filename.casefold(),
        )
        original_names = [item.filename for item in source_docs]
    elif source_mode == "upload":
        if not upload_items:
            raise ValueError("Добавьте хотя бы один DOCX.")
        original_names = [item[0] for item in upload_items]
    else:
        raise ValueError("Выберите копирование или загрузку документов.")

    requested_name = str(name or "").strip()
    if variant:
        resolved_name = safe_kit_name(variant)
    elif requested_name:
        resolved_name = safe_kit_name(requested_name)
    elif source_kit is not None:
        resolved_name = safe_kit_name(str(source_kit.get("release_clean") or source_kit.get("category") or category))
    else:
        resolved_name = safe_kit_name(category.replace("_", " "))
    name = resolved_name

    release_full = f"{name}({ke})"
    target_relative = PurePosixPath(category, release_full).as_posix()
    target_kit_id = kit_id_for_relative_dir(target_relative)
    if any(entry["kit_id"] == target_kit_id for entry in list_kit_models(root)):
        raise CreationDraftConflict("Комплект с таким названием и КЭ уже существует.")

    names = _ensure_unique_names(target_names or original_names)
    if len(names) != len(original_names):
        raise ValueError("Для каждого документа должно быть указано одно итоговое имя.")
    draft_uuid, directory = _new_draft({
        "operation": "create_kit", "category": category, "name": name, "ke": ke,
        "variant": variant, "aliases": aliases, "target_relative_dir": target_relative,
        "target_display_name": release_full,
        "target_kit_id": target_kit_id, "source_mode": source_mode,
        "source_kit_id": source_kit_id if source_kit else "", "files": [],
    })
    try:
        file_models = []
        pairs = zip(names, source_docs if source_mode == "copy" else upload_items)
        for index, (target_name, source) in enumerate(pairs):
            path = directory / f"document-{index}.docx"
            if source_mode == "copy":
                payload = source.path.read_bytes()
                if len(payload) > MAX_UPLOAD_BYTES:
                    raise CandidateUploadTooLarge("File exceeds the 10 MiB limit")
                atomic_write_bytes(path, payload)
                size, digest = len(payload), hashlib.sha256(payload).hexdigest()
                original_name = source.filename
            else:
                original_name, stream = source
                size, digest = _write_stream(stream, path)
            file_models.append({
                "index": index,
                "source_filename": Path(original_name).name,
                "target_filename": target_name,
                "size": size,
                "sha256": digest,
                "origin": source_mode,
            })
        metadata = read_json(directory / "metadata.json")
        metadata["files"] = file_models
        atomic_write_json(directory / "metadata.json", metadata)
        return validate_creation_draft(draft_uuid)
    except Exception:
        shutil.rmtree(directory, ignore_errors=True)
        raise


def create_document_draft(*, root: Path, kit_id: str, filename: str, stream: BinaryIO) -> dict[str, Any]:
    kit = resolve_kit(kit_id, root)
    if kit is None:
        raise ValueError("Комплект не найден.")
    filename = safe_docx_filename(filename)
    existing = [path.name.casefold() for path in Path(kit["path"]).glob("*.docx")]
    if filename.casefold() in existing:
        raise CreationDraftConflict("Документ с таким именем уже существует. Используйте действие «Заменить».")
    draft_uuid, directory = _new_draft({
        "operation": "add_document", "target_kit_id": kit_id,
        "target_relative_dir": kit["relative_dir"],
        "target_display_name": kit.get("release_full") or "",
        "files": [],
    })
    try:
        path = directory / "document-0.docx"
        size, digest = _write_stream(stream, path)
        metadata = read_json(directory / "metadata.json")
        metadata["files"] = [{
            "index": 0,
            "source_filename": filename,
            "target_filename": filename,
            "size": size,
            "sha256": digest,
            "origin": "upload",
        }]
        atomic_write_json(directory / "metadata.json", metadata)
        return validate_creation_draft(draft_uuid)
    except Exception:
        shutil.rmtree(directory, ignore_errors=True)
        raise


def get_creation_draft(draft_uuid: str) -> dict[str, Any]:
    directory = _draft_directory(draft_uuid)
    metadata = read_json(directory / "metadata.json", missing=None)
    if not isinstance(metadata, dict) or metadata.get("draft_uuid") != draft_uuid:
        raise CreationDraftNotFound("Draft not found")
    return metadata


def _validate_new_document(
    path: Path,
    *,
    target_relative_dir: str,
    target_filename: str,
    test_output: Path,
) -> dict[str, Any]:
    errors, external_images = inspect_docx(path)
    warnings: list[ValidationFailure] = []
    contract: dict[str, Any] = {}
    if external_images:
        warnings.append(ValidationFailure(
            "external_image_not_loaded",
            "Документ содержит внешние изображения; Центр шаблонов не загружает их автоматически.",
            "security",
        ))
    if not errors:
        try:
            contract = build_contract(path)
            unknown = sorted(
                {
                    str(value)
                    for value in contract.get("placeholders") or {}
                    if _placeholder_name(value) not in LEGACY_PLACEHOLDERS
                }
            )
            if unknown:
                errors.append(ValidationFailure("unknown_placeholder", "Найдены неподдерживаемые служебные ключи: " + ", ".join(unknown), "placeholders"))
            if contract.get("areas", {}).get("unsupported"):
                errors.append(ValidationFailure("placeholder_unsupported_area", "Служебные поля в колонтитулах и специальных областях не поддерживаются.", "placeholders"))
            if contract.get("damaged_placeholders"):
                errors.append(ValidationFailure("placeholder_whitespace_damaged", "Распознанное служебное поле повреждено пробелами.", "placeholders"))
            for signature in contract.get("jira_signatures") or []:
                if not signature.get("modern_supported"):
                    errors.append(ValidationFailure("jira_table_columns", "Новая Jira-таблица должна иметь поддерживаемые колонки ЗНИ/JIRA ID, описание задачи и тип.", "jira"))
            if "14061745" in PurePosixPath(target_relative_dir).name:
                for message in validate_ai_planner_template_structure(Document(path), template_name=target_filename):
                    errors.append(ValidationFailure("ai_planner_generator_contract", message, "generation"))
        except Exception:
            errors.append(ValidationFailure("contract_read", "Не удалось проверить структуру служебных полей.", "structure"))
    if not errors:
        generated, generation_errors = generate_synthetic_document(path, test_output)
        if generated is None:
            errors.extend(ValidationFailure(**item) for item in generation_errors)
    return {
        "ok": not errors,
        "errors": [item.as_dict() for item in errors],
        "warnings": [item.as_dict() for item in warnings],
        "contract": contract,
    }


def validate_creation_draft(draft_uuid: str) -> dict[str, Any]:
    with _draft_lock(draft_uuid):
        metadata = get_creation_draft(draft_uuid)
        if metadata.get("state") not in {"uploaded", "invalid", "valid"}:
            raise CreationDraftConflict("Черновик находится в несовместимом состоянии.")
        results = []
        for item in metadata.get("files") or []:
            index = int(item["index"])
            result = _validate_new_document(
                _draft_directory(draft_uuid) / f"document-{index}.docx",
                target_relative_dir=str(metadata.get("target_relative_dir") or ""),
                target_filename=str(item.get("target_filename") or ""),
                test_output=_draft_directory(draft_uuid) / f"test-{index}.docx",
            )
            results.append({**item, "validation": result})
        metadata["files"] = results
        metadata["validation"] = {"ok": bool(results) and all(item["validation"]["ok"] for item in results)}
        metadata["state"] = "valid" if metadata["validation"]["ok"] else "invalid"
        metadata["updated_at"] = utc_now()
        atomic_write_json(_draft_directory(draft_uuid) / "metadata.json", metadata)
        return metadata


def creation_draft_payload(draft_uuid: str, index: int, *, test: bool = False) -> tuple[bytes, dict[str, Any]]:
    metadata = get_creation_draft(draft_uuid)
    files = [item for item in metadata.get("files") or [] if int(item.get("index", -1)) == int(index)]
    if not files:
        raise CreationDraftNotFound("File not found")
    path = _draft_directory(draft_uuid) / (f"test-{index}.docx" if test else f"document-{index}.docx")
    if test and not files[0].get("validation", {}).get("ok"):
        raise CreationDraftConflict("Test document unavailable")
    try:
        return path.read_bytes(), files[0]
    except OSError as exc:
        raise CreationDraftNotFound("File not found") from exc


def publish_creation_draft(draft_uuid: str, root: Path) -> dict[str, Any]:
    root = Path(root)
    with _draft_lock(draft_uuid):
        metadata = get_creation_draft(draft_uuid)
        if metadata.get("state") != "valid":
            raise CreationDraftConflict("Черновик не готов к публикации.")
        kit_id = str(metadata.get("target_kit_id") or "")
        with _kit_lock(kit_id):
            relative = PurePosixPath(str(metadata["target_relative_dir"]))
            target_dir = root.joinpath(*relative.parts)
            files = list(metadata.get("files") or [])
            if metadata["operation"] == "create_kit":
                if target_dir.exists():
                    raise CreationDraftConflict("Комплект уже существует.")
                category_dir = target_dir.parent
                if not category_dir.is_dir() or category_dir.is_symlink():
                    raise CreationDraftConflict("Категория больше не доступна.")
                temporary = category_dir / f".oplot-kit-{draft_uuid}"
                temporary.mkdir(exist_ok=False)
                try:
                    for item in files:
                        shutil.copyfile(_draft_directory(draft_uuid) / f"document-{item['index']}.docx", temporary / item["target_filename"])
                    manifest = {
                        "name": metadata["name"], "ke": metadata["ke"],
                        "variant": metadata.get("variant") or "", "aliases": metadata.get("aliases") or [],
                    }
                    atomic_write_json(temporary / "manifest.json", manifest)
                    os.replace(temporary, target_dir)
                except Exception:
                    shutil.rmtree(temporary, ignore_errors=True)
                    raise
            else:
                kit = resolve_kit(kit_id, root)
                if kit is None:
                    raise CreationDraftConflict("Комплект больше не существует.")
                item = files[0]
                target = Path(kit["path"]) / item["target_filename"]
                if any(path.name.casefold() == target.name.casefold() for path in Path(kit["path"]).glob("*.docx")):
                    raise CreationDraftConflict("Документ с таким именем уже существует.")
                atomic_write_bytes(target, (_draft_directory(draft_uuid) / "document-0.docx").read_bytes())
        metadata["state"] = "published"
        metadata["updated_at"] = utc_now()
        atomic_write_json(_draft_directory(draft_uuid) / "metadata.json", metadata)
        clear_template_catalog_cache()
        clear_document_template_read_cache(root)
        return metadata


def cancel_creation_draft(draft_uuid: str) -> None:
    with _draft_lock(draft_uuid):
        metadata = get_creation_draft(draft_uuid)
        if metadata.get("state") == "published":
            raise CreationDraftConflict("Опубликованный черновик нельзя отменить.")
        shutil.rmtree(_draft_directory(draft_uuid), ignore_errors=False)


def document_can_be_deleted(document: ResolvedDocument) -> tuple[bool, str]:
    if list_history(document.document_id):
        return False, "У документа есть история версий. Удаление действующего документа заблокировано."
    active = [item for item in list_candidates(document.document_id) if item.get("state") not in {"published", "cancelled", "expired", "recovered", "conflict"}]
    if active:
        return False, "У документа есть активная новая версия. Сначала завершите или отмените её."
    return True, ""


def delete_active_document(document: ResolvedDocument) -> None:
    allowed, reason = document_can_be_deleted(document)
    if not allowed:
        raise CreationDraftConflict(reason)
    kit_id = kit_id_for_relative_dir(document.relative_directory)
    with _kit_lock(kit_id):
        siblings = list(document.path.parent.glob("*.docx"))
        if len(siblings) <= 1:
            raise CreationDraftConflict("Последний документ удаляется только вместе с комплектом.")
        document.path.unlink()
    clear_template_catalog_cache()
    clear_document_template_read_cache(document.template_root)


def kit_can_be_deleted(kit: dict[str, Any], root: Path) -> tuple[bool, str, list[ResolvedDocument]]:
    whitelist = build_document_whitelist(root, fresh=True)
    documents = [item for item in whitelist.values() if item.relative_directory.casefold() == str(kit["relative_dir"]).casefold()]
    if not documents:
        return False, "Комплект не содержит доступных документов.", []
    return True, "", documents


def _remove_runtime_tree(path: Path) -> None:
    if path.is_symlink():
        path.unlink(missing_ok=True)
    elif path.exists():
        shutil.rmtree(path)


def _creation_drafts_for_kit(kit_id: str) -> list[Path]:
    root = cache_root() / "creation-drafts"
    if not root.is_dir():
        return []
    result: list[Path] = []
    for metadata_path in root.glob("*/metadata.json"):
        metadata = read_json(metadata_path, missing=None)
        if isinstance(metadata, dict) and metadata.get("target_kit_id") == kit_id:
            result.append(metadata_path.parent)
    return result


def delete_active_kit(kit: dict[str, Any], root: Path) -> None:
    allowed, reason, documents = kit_can_be_deleted(kit, root)
    if not allowed:
        raise CreationDraftConflict(reason)
    with _kit_lock(kit["kit_id"]):
        # Deleting a complete kit is intentionally cascading: history and
        # unfinished replacements cannot be useful once their active files no
        # longer exist. Remove the visible kit first, then its opaque runtime
        # records so a partially cleaned operation never leaves active DOCX.
        shutil.rmtree(Path(kit["path"]))
        for document in documents:
            for candidate in list_candidates(document.document_id, allow_expired=True):
                _remove_runtime_tree(candidate_directory(str(candidate["candidate_uuid"])))
            _remove_runtime_tree(data_root() / "history" / document.document_id)
        for draft_directory in _creation_drafts_for_kit(str(kit["kit_id"])):
            _remove_runtime_tree(draft_directory)
    clear_template_catalog_cache()
    clear_document_template_read_cache(root)
